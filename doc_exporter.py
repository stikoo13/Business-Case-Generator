from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io

def export_to_word(inputs, financials, risks, ai_narrative):
    doc = Document()

    title = doc.add_heading(f"Business Case: {inputs['solution']}", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"Prepared for: {inputs['company']}")
    doc.add_paragraph(f"Date: {inputs['date']}")
    doc.add_paragraph("_" * 60)

    doc.add_heading("Financial Summary", level=1)
    table = doc.add_table(rows=4, cols=2)
    table.style = 'Table Grid'
    data = [
        ("Total ROI", f"{financials['roi_percent']}%"),
        ("Net Present Value (NPV)", f"{financials['npv']:,}"),
        ("Payback Period", f"{financials['payback_years']} years"),
        ("Implementation Cost", f"{inputs['cost']:,}"),
    ]
    for i, (label, value) in enumerate(data):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = value

    doc.add_paragraph("")

    for line in ai_narrative.split('\n'):
        if line.startswith('## '):
            doc.add_heading(line.replace('## ', ''), level=1)
        elif line.strip():
            doc.add_paragraph(line)

    doc.add_heading("Risk Matrix", level=1)
    risk_table = doc.add_table(rows=len(risks)+1, cols=4)
    risk_table.style = 'Table Grid'
    headers = ['Risk', 'Likelihood', 'Impact', 'Mitigation']
    for i, h in enumerate(headers):
        risk_table.rows[0].cells[i].text = h
    for i, risk in enumerate(risks):
        row = risk_table.rows[i+1]
        row.cells[0].text = risk['risk']
        row.cells[1].text = risk['likelihood']
        row.cells[2].text = risk['impact']
        row.cells[3].text = risk['mitigation']

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def export_to_pdf(inputs, financials, risks, ai_narrative):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                            rightMargin=inch, leftMargin=inch,
                            topMargin=inch, bottomMargin=inch)

    styles = getSampleStyleSheet()
    story = []

    # Title
    title_style = ParagraphStyle('Title', parent=styles['Title'], fontSize=18, spaceAfter=6)
    story.append(Paragraph(f"Business Case: {inputs['solution']}", title_style))
    story.append(Spacer(1, 6))
    story.append(Paragraph(f"Prepared for: {inputs['company']}", styles['Normal']))
    story.append(Paragraph(f"Date: {inputs['date']}", styles['Normal']))
    story.append(Spacer(1, 12))

    # Financial summary table
    story.append(Paragraph("Financial Summary", styles['Heading1']))
    fin_data = [
        ["Metric", "Value"],
        ["Total ROI", f"{financials['roi_percent']}%"],
        ["Net Present Value", f"{financials['npv']:,}"],
        ["Payback Period", f"{financials['payback_years']} years"],
        ["Implementation Cost", f"{inputs['cost']:,}"],
    ]
    fin_table = Table(fin_data, colWidths=[3*inch, 3*inch])
    fin_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#2c3e50')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f2f2f2')]),
        ('PADDING', (0,0), (-1,-1), 8),
    ]))
    story.append(fin_table)
    story.append(Spacer(1, 16))

    # AI narrative
    for line in ai_narrative.split('\n'):
        if line.startswith('## '):
            story.append(Paragraph(line.replace('## ', ''), styles['Heading1']))
        elif line.strip():
            story.append(Paragraph(line, styles['Normal']))
            story.append(Spacer(1, 4))

    story.append(Spacer(1, 12))

    # Risk matrix
    story.append(Paragraph("Risk Matrix", styles['Heading1']))
    risk_data = [["Risk", "Likelihood", "Impact", "Mitigation"]]
    for r in risks:
        risk_data.append([r['risk'], r['likelihood'], r['impact'], r['mitigation']])
    risk_table = Table(risk_data, colWidths=[1.5*inch, 1*inch, 1*inch, 3*inch])
    risk_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#2c3e50')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f2f2f2')]),
        ('PADDING', (0,0), (-1,-1), 6),
        ('FONTSIZE', (0,0), (-1,-1), 8),
        ('WORDWRAP', (0,0), (-1,-1), True),
    ]))
    story.append(risk_table)

    doc.build(story)
    buffer.seek(0)
    return buffer